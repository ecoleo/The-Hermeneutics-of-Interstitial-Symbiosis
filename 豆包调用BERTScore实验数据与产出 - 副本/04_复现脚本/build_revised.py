# -*- coding: utf-8 -*-
import os, copy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from content_revision import REPLACE, INS_AFTER_41

D = r"F:\文学数字人文阐释实践\故事新编研究案例\第四节 生成-实验路径-《故事新编》人机共生续写"
fn = [n for n in os.listdir(D) if n.startswith('第四节') and n.endswith('.docx') and not n.startswith('~$') and '修订版' not in n][0]
doc = Document(os.path.join(D, fn))
paras = list(doc.paragraphs)

def set_run_font(run, size=12, east='宋体', bold=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), east)

def replace_text(p, text):
    # 保留首个 run 的 rPr，删除其余 run
    runs = p.runs
    if runs:
        first = runs[0]
        first.text = text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        r = p.add_run(text); set_run_font(r)

# 1) 替换段落
for idx, txt in REPLACE.items():
    replace_text(paras[idx], txt)
print('replaced', len(REPLACE), 'paragraphs')

# 2) 在 paras[41] 后依次插入块
template_pPr = paras[37]._p.find(qn('w:pPr'))  # 正文段落属性模板

def new_para(text, kind='body'):
    p_el = OxmlElement('w:p')
    if kind == 'body' and template_pPr is not None:
        p_el.append(copy.deepcopy(template_pPr))
    p = Paragraph(p_el, paras[41]._parent)
    if kind == 'cap' or kind == 'tabtitle':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    if kind == 'cap':
        set_run_font(r, 10.5, '宋体')
    elif kind == 'tabtitle':
        set_run_font(r, 10.5, '黑体', bold=True)
    else:
        set_run_font(r, 12, '宋体')
    return p

def new_image_para(path):
    p_el = OxmlElement('w:p')
    p = Paragraph(p_el, paras[41]._parent)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(path, width=Cm(14.5))
    return p

cursor = paras[41]._p
tbl_holder = []
for kind, payload in INS_AFTER_41:
    if kind == 'p':
        el = new_para(payload, 'body')._p
    elif kind == 'img':
        el = new_image_para(payload)._p
    elif kind == 'cap':
        el = new_para(payload, 'cap')._p
    elif kind == 'tabtitle':
        el = new_para(payload, 'tabtitle')._p
    elif kind == 'table':
        data = payload
        t = doc.add_table(rows=len(data), cols=len(data[0]))
        # 手工加全框线
        tblPr = t._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            e = OxmlElement(f'w:{edge}')
            e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4')
            e.set(qn('w:space'),'0'); e.set(qn('w:color'),'000000')
            borders.append(e)
        tblPr.append(borders)
        t.alignment = 1  # center
        for i, row in enumerate(data):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.text = ''
                pp = cell.paragraphs[0]
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = pp.add_run(val)
                set_run_font(rr, 10.5, '宋体', bold=(i == 0))
        el = t._tbl
    cursor.addnext(el)
    cursor = el
    if kind == 'table':
        # 表格后补一个空段，避免与后续表格/段落粘连
        sp = OxmlElement('w:p'); cursor.addnext(sp); cursor = sp

out = "第四节 生成-实验型进路：生成式AI续写实验与意义归属的多重界定（学术审订修订版）.docx"
doc.save(out)
print('saved', out)
