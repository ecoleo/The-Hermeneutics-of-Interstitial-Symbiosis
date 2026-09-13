# -*- coding: utf-8 -*-
"""
修订稿二次修补（2026-09-13）：
1. 替换遗留的失真工具链句（八爪鱼/Gephi/Pandas → 真实工具链）；
2. 段[11](三) 抖音方法句：受阻表述 → 豆包快照数据表述；
3. 段[17] 抖音分析段开头：删除"不作定量断言"，回填豆包快照实测值；
4. 在方法段后插入两段新段落：WorkBuddy选用理由 + 交叉验证说明。
每处替换须唯一命中；插入段落复制锚点段 pPr 以继承正文格式。
"""
import copy, os, sys, zipfile, re, html
import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DST = os.path.join(ROOT, "交付_2026-09-12", "01_论文修订稿",
                   "第五节 接受-批判进路：《故事新编》数字传播生态的算法规训批判（修订稿）.docx")

R = [
    ("各平台公开数据经八爪鱼采集器与Python爬虫抓取，标签共现结构由Gephi构建关联网络，频次与偏差分布由Pandas统计并以Matplotlib可视化，全程仅处理平台前端可观测材料。",
     "各平台公开数据由WorkBuddy智能体工作平台统一调度采集与分析：B站与豆瓣材料经Python标准库脚本抓取（B站搜索接口的WBI签名由脚本自行实现），微信读书材料经平台官方授权接口（Agent API Gateway）获取，抖音材料由豆包智能体经人工辅助的受控浏览器采集，频次与偏差分布由Python统计脚本计算并以Matplotlib可视化，全程仅处理平台前端可观测材料，采集脚本、参数设置与原始数据均留存可复核。"),
    ("（三）短视频平台——原定以抖音“鲁迅说”话题为对象，但平台对网页端搜索实施验证码拦截、对接口调用实施签名校验，在“公开可观测的前端材料”边界内无法获得可核验样本，登录态网页搜索、web接口与第三方云采集模板市场三条路径均已尝试而不可得，故本节对短视频平台不作定量断言，仅作文献层面的机制性讨论；",
     "（三）短视频平台的内容分类与推荐话题——抖音材料由豆包智能体经人工辅助的受控浏览器采集（研究者本人两次手动通过登录与图形验证码，未绕过平台机制），包括“鲁迅说”话题快照147条（其中相关样本136条）与“故事新编 鲁迅”检索快照18条（采集时点2026年9月12日），据此分析内容形态构成与传播特征；该数据为可见性快照而非概率抽样，其科学适用性见本节交叉验证说明；"),
    ("短视频平台集中体现了碎片化传播的机制性力量。须先作方法学说明：本节原定以抖音“鲁迅说”话题为考察对象，但平台的反爬机制（网页端搜索返回验证码中间页、接口调用要求签名校验）使可核验样本在“公开可观测的前端材料”边界内不可得，第三方云采集模板市场亦无可用抖音模板，故以下讨论不作定量断言，仅基于既有文献作机制性分析，相关比例判断留待数据可得时验证。",
     "短视频平台集中体现了碎片化传播的机制性力量。须先作方法学说明：抖音材料由豆包智能体经人工辅助的受控浏览器采集，为单账号、单时点、综合排序下的可见性快照（2026年9月12日），本节据此只作“可见性结构描述”，不作总体推断。快照显示，“鲁迅说”话题（n=147，其中相关样本136）的内容构成以碎片化形态为主导：名句呈现占41.5%、仿写玩梗占18.4%、情绪共鸣占10.2%，相关样本中碎片化形态合计75.7%，剧情演绎为0；点赞中位数576，38条点赞过万，最高达274.6万。与之形成鲜明对照，“故事新编 鲁迅”检索入口（n=18）以知识讲解为主（83.3%），深度内容合计94.4%，点赞中位数仅81.5。同一平台内部，泛话题推荐入口与篇目词检索入口呈现“碎片化—深度化”的梯度分化：算法的可见性分配因入口性质而异，主动检索入口仍保留深度内容的可达性。"),
]

INSERT_A = ("本研究选用WorkBuddy智能体工作平台执行数据采集与分析，理由有三。其一，工具链的可审计性："
            "全部采集与统计均以可读脚本形式沉淀，检索词、排序条件、翻页上限、请求间隔、词典词表等参数"
            "在分析前显式固定并随数据一并存档，任何结论均可经脚本复跑检验，满足数字人文研究对可复现性的要求。"
            "其二，数据来源的合规与透明：微信读书材料经平台官方授权接口（Agent API Gateway）获取，"
            "B站与豆瓣仅采集前端公开可观测材料，登录凭据仅经环境变量临时传入、不落入任何数据文件，"
            "与本文“公开可观测材料”的操作边界及研究伦理一致。其三，统计方法的常规性与可核验性："
            "类别测量采用分析前冻结的预登记词典（防止事后挑选词表迎合结论），比例估计采用Wilson 95%置信区间，"
            "组间差异采用两比例z检验，分词采用jieba与TF-IDF，均为教科书级的常规方法，无黑箱模型参与；"
            "启发式情感指标（如SnowNLP输出）按既定数据纪律不进入正文。在此人机分工中，"
            "智能体承担可形式化的采集与计算，研究者保留词典设计、口径裁定与阐释判断——"
            "这一分工本身即是本文“居间共生阐释”方法论在研究实践层面的贯彻。")

INSERT_B = ("为保证上述分析的科学性与真实性，本研究引入豆包智能体进行独立交叉验证，两智能体的数据能力呈互补关系："
            "WorkBuddy在“公开可观测前端材料”的严格边界内无法获得抖音数据（网页端验证码拦截、接口签名校验、"
            "第三方云采集模板缺位），而豆包智能体与抖音同属字节跳动生态，具备受控浏览器自动化能力，"
            "经研究者本人两次人工接管（扫码登录、图形验证码）完成抖音快照采集；须强调，两次人机验证均由研究者本人完成，"
            "未绕过平台机制，且豆包所获同样是单账号、单时点、综合排序下的可见性快照，而非概率抽样。"
            "交叉验证结果显示两智能体分析具有高度一致性：B站数据为独立实现的逐值复现——双方各自实现WBI签名与双轨采样，"
            "均得797条（230/279/288），类别占比与检验值完全一致（文学品质80.0%→96.7%，有声说书15.7%→70.0%，z=−8.32）；"
            "豆瓣数据口径收敛——豆包200条均分4.58、五星65.3%，WorkBuddy扩样497条均分4.59、五星67.5%，"
            "方向与量级一致（差异源于采样框不同：前者含热门与最新两轨，后者为热门轨扩样）；"
            "微信读书数据呈超集关系——豆包20条点评与20条热门划线为WorkBuddy 1036条用户文本与400条章节划线的子集，"
            "篇目分布相容；抖音数据则由豆包单方提供，填补平台矩阵的空白。"
            "可见性快照方法对本研究是否科学可用，取决于研究目的：本研究的对象不是“平台总体内容”，"
            "而是“算法排序下读者实际可见的内容结构”——可见性本身即研究问题，快照方法因此不是次优妥协，"
            "而是与研究目的同构的观测设计。其科学可用性的三重条件本研究均已满足：快照条件（账号、时点、排序、检索入口）"
            "全部显式记录；结论口径限定为“可见性结构描述”，因果措辞统一降级为“与算法排序机制相一致的可见性再分配”，"
            "不作总体推断；关键结论经双智能体独立实现交叉验证。同时须明确其失效边界：凡涉及“平台总体占比”"
            "“用户整体偏好”“历时趋势”的推断，快照方法均不支持，本文亦未作此类断言。")

ANCHOR = "以预登记词典做主题结构分析；采集时点2026年9月12日。"


def replace_in_paragraph(p, old, new):
    full = "".join(r.text for r in p.runs)
    if old not in full:
        return False
    p.runs[0].text = full.replace(old, new)
    for r in p.runs[1:]:
        r.text = ""
    return True


def insert_after(anchor_p, text):
    new_p = copy.deepcopy(anchor_p._p)
    # 清空 deepcopy 带来的所有 run，仅保留 pPr
    for child in list(new_p):
        if not child.tag.endswith("}pPr"):
            new_p.remove(child)
    anchor_p._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    np = Paragraph(new_p, anchor_p._parent)
    # 复制锚点首个 run 的 rPr 保证字体一致
    r = np.add_run(text)
    if anchor_p.runs:
        src_rpr = anchor_p.runs[0]._element.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
        if src_rpr is not None:
            r._element.insert(0, copy.deepcopy(src_rpr))
    return np


def main():
    doc = docx.Document(DST)
    misses = []
    for i, (old, new) in enumerate(R, 1):
        found = 0
        for p in doc.paragraphs:
            if old in "".join(r.text for r in p.runs):
                found += 1
                replace_in_paragraph(p, old, new)
        if found != 1:
            misses.append((i, found, old[:36]))
    if misses:
        print("替换未唯一命中，中止：")
        for m in misses:
            print("  ", m)
        sys.exit(1)
    print(f"替换 {len(R)}/{len(R)} 处唯一命中")

    anchor = None
    for p in doc.paragraphs:
        if ANCHOR in "".join(r.text for r in p.runs):
            anchor = p
    if anchor is None:
        print("插入锚点未找到"); sys.exit(1)
    # 先插 B 再插 A（addnext 逆序），使顺序为 锚点 → A → B
    insert_after(anchor, INSERT_B)
    insert_after(anchor, INSERT_A)
    doc.save(DST)
    print("已插入 2 个新段落（WorkBuddy选用理由、交叉验证说明）")

    xml = zipfile.ZipFile(DST).read("word/document.xml").decode("utf-8")
    text = html.unescape("".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml, re.S)))
    checks = {
        "工具链新句": "各平台公开数据由WorkBuddy智能体工作平台统一调度采集与分析" in text,
        "八爪鱼句已灭": "八爪鱼采集器与Python爬虫抓取" not in text,
        "Gephi已灭": "Gephi" not in text,
        "段11三新句": "“鲁迅说”话题快照147条" in text,
        "段17实测": "相关样本中碎片化形态合计75.7%" in text,
        "不作断言旧句已灭": "不作定量断言" not in text,
        "插入A": "本研究选用WorkBuddy智能体工作平台执行数据采集与分析" in text,
        "插入B": "本研究引入豆包智能体进行独立交叉验证" in text,
    }
    for k, v in checks.items():
        print(f"  {'✔' if v else '✘'} {k}")
    print("图片数:", xml.count("<pic:pic"), "总字数:", len(text))


if __name__ == "__main__":
    main()
