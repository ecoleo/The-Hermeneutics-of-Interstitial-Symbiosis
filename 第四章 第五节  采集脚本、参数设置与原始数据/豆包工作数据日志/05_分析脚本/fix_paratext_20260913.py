# -*- coding: utf-8 -*-
"""2026-09-13（三次）：方法段补入"数字副文本"概念界定。
1) 首句由"运用数字副文本分析方法"改为"以数字副文本为研究对象，运用内容分析与描述统计方法"（研究对象≠方法的分层）；
2) 在"分析对象限定……操作边界诚实。"之后插入概念界定句（热奈特界定 + 数字环境延伸 + 分层说明）。
增量编辑、继承样式，仅改方法段。"""
import copy
import os

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = r'F:\文学数字人文阐释实践\故事新编研究案例\第五节《故事新编》数字传播生态的算法规训批判'
DST = os.path.join(BASE, '第五节_数字副文本采集与修订交付_20260912', '03_论文修订稿',
                   '第五节 接受-批判进路：《故事新编》数字传播生态的算法规训批判（数字副文本实证修订稿）.docx')

OLD_HEAD = ('为了具体揭示算法系统如何塑造《故事新编》的接受框架，本部分运用数字副文本分析方法，'
            '依托豆包工作智能体的内嵌数据分析工具，系统收集各主要数字平台上《故事新编》相关内容的可观测材料。')
NEW_HEAD = ('为了具体揭示算法系统如何塑造《故事新编》的接受框架，本部分以数字副文本为研究对象，'
            '依托豆包工作智能体的内嵌数据分析工具，运用内容分析与描述统计方法，'
            '系统收集各主要数字平台上《故事新编》相关内容的可观测材料。')

OLD_SCOPE = '分析对象限定为公开可获取的数字副文本，不涉及商业平台后台数据，以保持语境算法进路的操作边界诚实。'
INSERT = ('这里所称“数字副文本”，沿用热奈特关于“副文本是环绕正文并引导读者进入文本的门槛装置”的界定，'
          '并延伸到数字环境：印刷时代由作者与出版者主导的书名、序言、注释、评论等，'
          '在数字平台演变为由平台算法与用户共同生产的标签、分类、话题、短评、划线与人气排序。'
          '需要说明的是，“数字副文本”在此指研究对象而非某种既定分析程式，'
          '本节实际采用的分析方法是内容分析与描述统计（预登记词典、开放编码、频次与占比、比例检验），'
          '以保证每一条归类可人工复核、全部统计可复现。')
NEW_SCOPE = OLD_SCOPE + INSERT


def set_font(run, tmpl_rpr, bold):
    if tmpl_rpr is not None:
        run._element.insert(0, copy.deepcopy(tmpl_rpr))
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    rf.set(qn('w:ascii'), 'Times New Roman')
    rf.set(qn('w:hAnsi'), 'Times New Roman')
    rf.set(qn('w:eastAsia'), '宋体')


def patch_para(p):
    """在方法段上应用两处修改；返回改动明细。"""
    text = p.text
    hits = []
    if OLD_HEAD in text:
        text = text.replace(OLD_HEAD, NEW_HEAD)
        hits.append('首句分层')
    if OLD_SCOPE in text:
        text = text.replace(OLD_SCOPE, NEW_SCOPE)
        hits.append('概念界定插入')
    if not hits:
        return hits
    bold_first = bool(p.runs and p.runs[0].font.bold)
    tmpl = None
    if p.runs and p.runs[0]._element.rPr is not None:
        tmpl = copy.deepcopy(p.runs[0]._element.rPr)
    for r in list(p._element.findall(qn('w:r'))):
        p._element.remove(r)
    if bold_first and '。' in text:
        i = text.index('。') + 1
        r1 = p.add_run(text[:i]); set_font(r1, tmpl, True)
        r2 = p.add_run(text[i:]); set_font(r2, tmpl, False)
    else:
        r = p.add_run(text); set_font(r, tmpl, False)
    return hits


d = docx.Document(DST)
meth = None
for p in d.paragraphs:
    if p.text.startswith('为了具体揭示算法系统如何塑造'):
        meth = p
        break
assert meth is not None, '未定位到方法段'

hits = patch_para(meth)
print('改动:', hits if hits else '无命中!')
d.save(DST)
print('已保存:', DST)

d2 = docx.Document(DST)
full = '\n'.join(p.text for p in d2.paragraphs)
print('回读段落数:', len(d2.paragraphs), '图片:', len(d2.inline_shapes))
print('“运用数字副文本分析方法”残留:', '运用数字副文本分析方法' in full)
print('“以数字副文本为研究对象”在文:', '以数字副文本为研究对象' in full)
print('概念界定句在文:', '这里所称“数字副文本”' in full)
