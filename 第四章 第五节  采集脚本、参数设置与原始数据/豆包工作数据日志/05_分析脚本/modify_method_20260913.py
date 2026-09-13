# -*- coding: utf-8 -*-
"""2026-09-13：方法部分二次修订。
1) 在方法段之前插入“方法论说明”段（论证选用豆包工作智能体及其内嵌数据分析工具的科学性）；
2) 替换方法段（保持原四类平台分析框架，工具链改为实际执行工具，保留实证采集口径）；
3) 在方法段之后插入“交叉验证说明”段（WorkBuddy 抖音受阻 vs 豆包浏览器补足；可见性快照方法可用性论证）。
增量编辑：TARGET 原地写回，继承原稿字体与段落格式，引导句加粗风格与正文一致。
"""
import copy
import os

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

BASE = r'F:\文学数字人文阐释实践\故事新编研究案例\第五节《故事新编》数字传播生态的算法规训批判'
DST = os.path.join(BASE, '第五节_数字副文本采集与修订交付_20260912', '03_论文修订稿',
                   '第五节 接受-批判进路：《故事新编》数字传播生态的算法规训批判（数字副文本实证修订稿）.docx')

RATIONALE = ('本研究的数据采集与分析依托豆包工作智能体的内嵌数据分析工具完成，选择理由有三。'
             '其一，执行边界的契合性。本研究的采集纪律要求仅处理平台前端可观测材料、不触碰任何后台数据、'
             '不绕过平台验证——抖音的登录与验证码均由研究者本人手动完成——豆包工作智能体将标准库爬虫、'
             '内置受控浏览器与平台官方开放接口纳入同一受控执行环境，使上述纪律成为可强制执行的硬约束而非事后声明。'
             '其二，数据来源的可追溯性。全部采集脚本逐条记录检索关键词、排序方式、翻页深度与采集时点，'
             '原始数据以JSONL逐行落盘，正文每一个统计量均由原始数据现算生成，均可回溯至具体记录与参数，'
             '不存在不可复现的中间环节。其三，统计方法的透明性。分析采用预先登记、词表冻结于正式计数之前的编码本'
             '（B站四类词典、抖音六类开放编码并经两轮人工校准），统计上使用Wilson 95%置信区间、两比例z检验'
             '与jieba分词词频等公开可复现方法，工具版本、阈值与全部参数在复现说明中逐项披露；'
             '原方案中的八爪鱼采集器与Gephi在本环境中不具备可执行条件，已分别以标准库请求与标签频率统计替代，'
             '全部结论不依赖任何特定商业软件。数据来源与统计方法由此满足可核验、可复现、口径一致三项要求，'
             '本节分析结果与正文论述的对应关系可在复现说明的一致性核对表中逐条查验，从而保证结论的科学性与可信度。')

METHOD_NEW = ('为了具体揭示算法系统如何塑造《故事新编》的接受框架，本部分运用数字副文本分析方法，'
              '依托豆包工作智能体的内嵌数据分析工具，系统收集各主要数字平台上《故事新编》相关内容的可观测材料。'
              '分析对象限定为公开可获取的数字副文本，不涉及商业平台后台数据，以保持语境算法进路的操作边界诚实。'
              '如第一节所述，商业平台的推荐算法属于不可获取的黑箱，任何声称可以“完全破解”算法的宣称在方法论上都是不诚实的；'
              '本节所能观测与统计的，仅是算法在前端“落地”后呈现的可见结果（标签、排序、划线人数、点赞量等），'
              '而非算法本身的权重与内部逻辑。数据于2026年9月11日至12日分平台采集，其中B站与豆瓣经自建Python爬虫脚本抓取，'
              '抖音经内置受控浏览器在登录态下采集，微信读书经平台官方开放接口获取；频次、占比与偏差分布由Pandas统计'
              '并以Matplotlib可视化，辅以jieba分词词频与Wilson置信区间、比例差异检验，全程仅处理平台前端可观测材料。'
              '分析对象包括四类平台的可观测材料：（一）B站视频的标签与分类——以“故事新编/鲁迅”为搜索关键词'
              '统计相关视频所使用的标签与分类，分析其频率分布与组合方式，并按“综合排序”与“最新排序”双轨采样'
              '以观察排序机制对可见性的再分配，按“视频×关键词”去重后共797条；（二）豆瓣的书目分类与读者标注——'
              '收集《故事新编》（书页编号2046909）页面的读者标注，以“热门”“最新”排序各100条共200条短评'
              '构建高频词分布（该书页“常用标签”由前端动态加载、无法通过公开途径获取，故以短评正文分词替代标签词云，'
              '页面声明短评总量8790条，样本覆盖率2.28%）；（三）短视频平台的内容分类与推荐话题——以抖音平台为例，'
              '在登录态浏览器中检索“鲁迅说”与“故事新编 鲁迅”，分析其话题归属与传播特征，滚动加载去重后'
              '分别获得147条与18条结果，并按预先登记、经两轮开放编码校准后锁定的编码本作内容形式分类；'
              '（四）微信读书的“读者说”摘录推荐内容——经平台官方开放接口收集推荐的读者划线句子与点评内容，'
              '热门划线全书共449条（服务端按热度返回前20条）、公开点评共547条（返回前20条），分析其主题分布。'
              '上述样本均为特定时点、特定检索条件下的“可见性快照”，而非平台内容总体的概率抽样；'
              '这一限度将在后文的因果措辞中始终保留。')

CROSSVAL = ('数据来源的效度进一步通过另一工作智能体（WorkBuddy）的独立采集加以交叉验证。'
            'WorkBuddy在2026年9月11日至12日的独立执行中完成了B站（N=797）、豆瓣（N=497，两轮合并去重）、'
            '微信读书（用户生成文本1036条、章节热门划线400条）的采集，但抖音采集的三条路径均受阻：'
            '网页端返回验证码中间页，搜索接口因缺少X-Bogus/msToken签名而不返回数据，云采集模板市场无抖音模板，'
            '故其修订稿对抖音不作定量断言，仅作文献层面机制性讨论。豆包工作智能体则借助内置受控浏览器，'
            '在研究者本人完成登录与验证码后，对抖音“鲁迅说”（n=147）与“故事新编 鲁迅”（n=18）完成采集，'
            '补足了前者缺失的短视频维度；须如实声明的是，该抖音数据为单账号、单时点、特定检索条件下的可见性快照，'
            '而非概率抽样。两个智能体独立采集、工具链不同（标准库请求与内置浏览器）、样本范围不同（如豆瓣200条与497条），'
            '在同一平台上的结论却高度趋同：豆瓣均分4.58与4.59、四至五星占比92.8%与92.5%几乎一致；'
            'B站均显示文学品质类标签占绝对多数、幽默解构类占比不超过8.1%；微信读书均未发现“毒舌化”主导的证据，'
            '热门划线均呈多元严肃分布。这种跨智能体、跨工具、跨样本范围的收敛，使各平台结论的真实性得到相互印证。'
            '就方法论而言，可见性快照对本节研究目的是科学可用的：本节研究的对象不是全体用户的传播效果，'
            '而是算法在给定检索条件下对可见性的再分配——快照恰是对这一对象的前端直接测量；'
            '在“故事新编 鲁迅”等长尾词下，前端结果近乎穷尽（n=18即该词的结果总量），更接近总体而非样本；'
            '且全部结论的措辞均限定于特定时点与检索条件下的“可见性结构”，不做普遍性外推，与规训批判的论证目标相符。')


def set_run_font(run, tmpl_rpr, bold, size_pt=12):
    if tmpl_rpr is not None:
        run._element.insert(0, copy.deepcopy(tmpl_rpr))
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    rf.set(qn('w:ascii'), 'Times New Roman')
    rf.set(qn('w:hAnsi'), 'Times New Roman')
    rf.set(qn('w:eastAsia'), '宋体')


def add_text(p, text, tmpl_rpr, bold_first):
    if bold_first and '。' in text:
        i = text.index('。') + 1
        r1 = p.add_run(text[:i]); set_run_font(r1, tmpl_rpr, True)
        r2 = p.add_run(text[i:]); set_run_font(r2, tmpl_rpr, False)
    else:
        r = p.add_run(text); set_run_font(r, tmpl_rpr, False)


def new_para_near(anchor_par, text, before=False):
    parent = anchor_par._parent
    el = OxmlElement('w:p')
    if before:
        anchor_par._element.addprevious(el)
    else:
        anchor_par._element.addnext(el)
    p = Paragraph(el, parent)
    # 继承锚点段段落格式（首行缩进、行距等）
    if anchor_par._element.pPr is not None:
        p._element.insert(0, copy.deepcopy(anchor_par._element.pPr))
    tmpl = None
    if anchor_par.runs and anchor_par.runs[0]._element.rPr is not None:
        tmpl = copy.deepcopy(anchor_par.runs[0]._element.rPr)
    add_text(p, text, tmpl, True)  # 引导句加粗，与正文既有引导段一致
    return p


def replace_para(p, text):
    bold_first = bool(p.runs and p.runs[0].font.bold)
    tmpl = None
    if p.runs and p.runs[0]._element.rPr is not None:
        tmpl = copy.deepcopy(p.runs[0]._element.rPr)
    for r in list(p._element.findall(qn('w:r'))):
        p._element.remove(r)
    add_text(p, text, tmpl, bold_first)


d = docx.Document(DST)

# 定位现有方法段（唯一锚点）
meth = None
for p in d.paragraphs:
    if p.text.startswith('为了具体揭示算法系统如何塑造'):
        meth = p
        break
assert meth is not None, '未定位到方法段'

# 1) 方法论说明：插在方法段之前
new_para_near(meth, RATIONALE, before=True)
# 2) 替换方法段
replace_para(meth, METHOD_NEW)
# 3) 交叉验证说明：插在方法段之后
new_para_near(meth, CROSSVAL, before=False)

d.save(DST)
print('已保存:', DST)

# 回读校验
d2 = docx.Document(DST)
print('回读 段落数:', len(d2.paragraphs), ' 内嵌图片:', len(d2.inline_shapes))
for i, p in enumerate(d2.paragraphs):
    t = p.text.strip()
    if t.startswith('本研究的数据采集与分析依托') or t.startswith('为了具体揭示算法') or t.startswith('数据来源的效度进一步') or (t.startswith('图') and len(t) < 80):
        print('  [%d] %s' % (i, t[:42]))
