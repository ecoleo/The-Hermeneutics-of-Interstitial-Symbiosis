# -*- coding: utf-8 -*-
"""在交付副本上执行实证修订：整段替换 + 嵌入图表，继承原稿字体与段落样式。"""
import copy
import os
import sys

import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from revision_texts import REPLACE, FIGURES

BASE = r'F:\文学数字人文阐释实践\故事新编研究案例\第五节《故事新编》数字传播生态的算法规训批判'
DST = os.path.join(BASE, '第五节_数字副文本采集与修订交付_20260912', '03_论文修订稿',
                   '第五节 接受-批判进路：《故事新编》数字传播生态的算法规训批判（数字副文本实证修订稿）.docx')
FIGDIR = os.path.join(BASE, '第五节_数字副文本采集与修订交付_20260912', '04_分析图表')

d = docx.Document(DST)


def set_run_font(run, tmpl_rpr, bold, size_pt=12):
    if tmpl_rpr is not None:
        run._element.insert(0, copy.deepcopy(tmpl_rpr))
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    rf.set(qn('w:ascii'), 'Times New Roman')
    rf.set(qn('w:hAnsi'), 'Times New Roman')
    rf.set(qn('w:eastAsia'), '宋体')


def replace_para(p, text):
    # 是否沿用“引导句加粗”：看原段首个 run
    bold_first = bool(p.runs and p.runs[0].font.bold)
    tmpl = None
    if p.runs and p.runs[0]._element.rPr is not None:
        tmpl = copy.deepcopy(p.runs[0]._element.rPr)
    for r in list(p._element.findall(qn('w:r'))):
        p._element.remove(r)
    if bold_first and '。' in text:
        i = text.index('。') + 1
        r1 = p.add_run(text[:i]); set_run_font(r1, tmpl, True)
        r2 = p.add_run(text[i:]); set_run_font(r2, tmpl, False)
    else:
        r = p.add_run(text); set_run_font(r, tmpl, False)


def new_para_after(anchor_el, parent):
    el = OxmlElement('w:p')
    anchor_el.addnext(el)
    return Paragraph(el, parent)


def insert_figure_after(anchor_par, img_name, caption):
    parent = anchor_par._parent
    # caption 先插，再插 img，保证 anchor -> img -> caption
    cap = new_para_after(anchor_par._element, parent)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)
    cr = cap.add_run(caption)
    cr.font.size = Pt(10.5); cr.font.bold = False
    rpr = cr._element.get_or_add_rPr()
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman')
    rf.set(qn('w:eastAsia'), '宋体'); rpr.insert(0, rf)

    imgp = new_para_after(anchor_par._element, parent)
    imgp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    imgp.paragraph_format.first_line_indent = Pt(0)
    imgp.add_run().add_picture(os.path.join(FIGDIR, img_name), width=Inches(5.8))
    return cap  # 下一张图接在图题之后


# 1) 文本替换（先按索引锁定对象）
targets = {idx: d.paragraphs[idx] for idx in REPLACE}
for idx, txt in REPLACE.items():
    replace_para(targets[idx], txt)
    print('替换 段%d (%d字)' % (idx, len(txt)))

# 2) 插图（同锚点多图按顺序串联）
last_anchor = None
cursor = None
for idx, img, cap in FIGURES:
    if idx != last_anchor:
        cursor = targets[idx] if idx in targets else d.paragraphs[idx]
        last_anchor = idx
    cursor = insert_figure_after(cursor, img, cap)
    print('插图 段%d <- %s' % (idx, img))

d.save(DST)
print('\n已保存:', DST)

# 3) 回读校验
d2 = docx.Document(DST)
print('回读 段落数:', len(d2.paragraphs), ' 内嵌图片:', len(d2.inline_shapes))
for i, p in enumerate(d2.paragraphs):
    t = p.text.strip()
    if t.startswith('图') and len(t) < 80:
        print('  图题@段%d:' % i, t[:60])
