# -*- coding: utf-8 -*-
"""对理论段“扁平化”做经验边界限定，使其与第二部分梯度结论一致。"""
import copy
import os
import docx
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = r'F:\文学数字人文阐释实践\故事新编研究案例\第五节《故事新编》数字传播生态的算法规训批判'
DST = os.path.join(BASE, '第五节_数字副文本采集与修订交付_20260912', '03_论文修订稿',
                   '第五节 接受-批判进路：《故事新编》数字传播生态的算法规训批判（数字副文本实证修订稿）.docx')
d = docx.Document(DST)

NEW = ("第一，意义深度的扁平化。如前文所分析，《故事新编》的“油滑”不是单纯的语言技巧，而是鲁迅在特定历史语境中"
"发展出的一种复杂的文体政治。它既是对古史辨派历史哲学的回应，也是对1930年代文化界论争的介入，更是鲁迅晚期思想中"
"“杂”之伦理的文学实践。在泛话题推荐这一流量最集中的入口，算法框架倾向于把这一复杂的文体政治压缩为“毒舌”与"
"“幽默”的标签，这构成对《故事新编》意义深度的扁平化——第二部分的实测同时表明，在主动检索与文学化社群入口，"
"文学品质类内容仍占多数，因而这种扁平化是流量逻辑作用下的选择性后果，而非平台内容的全部实况。这种扁平化不是对"
"鲁迅的简单“误读”，而是一种更根本的阐释性遮蔽：在被算法主导的那部分接触路径上，它取消了理解《故事新编》所必需的"
"历史意识和知识准备，将鲁迅变成一个不需要语境即可被即时消费的“金句生产者”。李宁在分析深度媒介化时代的故事危机时"
"指出，故事“越来越倾向于被处理为数字信息，讲述方式日趋单一”，传播受制于数据主义与流量逻辑。这一判断描述了《故事"
"新编》在泛推荐路径中的命运：复杂的文体政治被压缩为“数字信息”，多元讲述方式被简化为算法可识别的标签。崔波、朱咫渝"
"在比较抖音与B站荐书视频对阅读生态的重塑时发现，荐书视频“在流量的广度与思想的深度、商业的效率与文化的价值之间，"
"生长出充满张力却也孕育可能的复杂图景”。这种张力恰恰表明，扁平化不是算法推荐的必然结果，而是流量逻辑压倒文化逻辑"
"时的选择性后果，且其强度随入口与平台而变化。当商业效率的考量优先于文化价值的守护时，算法便倾向于将意义深度压缩为"
"可快速消费的标签。")

target = None
for p in d.paragraphs:
    if p.text.strip().startswith('第一，意义深度的扁平化'):
        target = p
        break
assert target is not None, '未找到目标段'

tmpl = copy.deepcopy(target.runs[0]._element.rPr) if target.runs and target.runs[0]._element.rPr is not None else None
for r in list(target._element.findall(qn('w:r'))):
    target._element.remove(r)
run = target.add_run(NEW)
if tmpl is not None:
    run._element.insert(0, copy.deepcopy(tmpl))
run.font.size = Pt(12); run.font.bold = False
rpr = run._element.get_or_add_rPr()
rf = rpr.find(qn('w:rFonts'))
if rf is None:
    rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman'); rf.set(qn('w:eastAsia'), '宋体')

d.save(DST)
print('已限定扁平化段，新字数', len(NEW))
d2 = docx.Document(DST)
print('回读段落:', len(d2.paragraphs), '图片:', len(d2.inline_shapes))
